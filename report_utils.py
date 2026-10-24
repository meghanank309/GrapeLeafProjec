from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


EXCEL_FILE = Path(__file__).parent / "grape_leaf_analysis_history.xlsx"

HEADERS = [
    "Analysis Date & Time",
    "Image Filename",
    "MobileNetV2 Prediction",
    "MobileNetV2 Score (%)",
    "VGG16 Prediction",
    "VGG16 Score (%)",
    "ResNet50 Prediction",
    "ResNet50 Score (%)",
    "EfficientNetB3 Prediction",
    "EfficientNetB3 Score (%)",
    "All Model Predictions",
]


def append_analysis_to_excel(filename, results):
    """Append one successful leaf analysis to the Excel history."""

    analysis_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    if EXCEL_FILE.exists():
        workbook = load_workbook(EXCEL_FILE)
        worksheet = workbook.active
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Analysis History"
        worksheet.append(HEADERS)

        # Format the header row
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="287D4A",
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

    # Store results in a predictable order
    expected_models = [
        "MobileNetV2",
        "VGG16",
        "ResNet50",
        "EfficientNetB3",
    ]

    row = [analysis_time, filename]
    all_predictions = []

    for model_name in expected_models:
        result = results.get(model_name)

        if result is None:
            row.extend(["Not available", None])
            all_predictions.append(f"{model_name}: Not available")
        else:
            prediction = result["predicted_class"]
            score = result["confidence"]

            row.extend([prediction, score])
            all_predictions.append(
                f"{model_name}: {prediction} ({score:.2f}%)"
            )

    row.append(" | ".join(all_predictions))
    worksheet.append(row)

    # Make the workbook easier to read
    for column_cells in worksheet.columns:
        column_letter = column_cells[0].column_letter
        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )
        worksheet.column_dimensions[column_letter].width = min(
            max(max_length + 2, 14),
            42,
        )

    workbook.save(EXCEL_FILE)
    return analysis_time

def create_pdf_report(filename, results, info, image, analysis_time):
    """Create a downloadable PDF report for one grape leaf analysis."""
    from io import BytesIO

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
        Image as PDFImage,
    )

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=42,
        leftMargin=42,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ReportTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            textColor=colors.HexColor("#287D4A"),
            fontSize=20,
            spaceAfter=12,
        )
    )

    story = []

    story.append(Paragraph("Grape Leaf Analysis Report", styles["ReportTitle"]))
    story.append(Paragraph(
        "AI Assisted Grape Leaf Disease Detection",
        styles["Heading2"],
    ))
    story.append(Spacer(1, 8))

    story.append(Paragraph(
        f"<b>Image filename:</b> {filename}",
        styles["BodyText"],
    ))
    story.append(Paragraph(
        f"<b>Analysis date and time:</b> {analysis_time}",
        styles["BodyText"],
    ))
    story.append(Spacer(1, 12))

    # Include the analyzed image when possible
    try:
        image_buffer = BytesIO()
        image.save(image_buffer, format="JPEG")
        image_buffer.seek(0)

        leaf_image = PDFImage(
            image_buffer,
            width=3.0 * inch,
            height=3.0 * inch,
            kind="proportional",
        )
        story.append(leaf_image)
        story.append(Spacer(1, 12))
    except Exception:
        story.append(Paragraph(
            "The analyzed image could not be embedded in this report.",
            styles["Italic"],
        ))

    story.append(Paragraph("Model Predictions", styles["Heading1"]))

    table_data = [
        ["Model", "Predicted Class", "Top Score (%)"]
    ]

    for model_name, result in results.items():
        table_data.append([
            str(model_name),
            str(result.get("predicted_class", "Unavailable")),
            f"{float(result.get('confidence', 0)):.2f}",
        ])

    prediction_table = Table(
        table_data,
        colWidths=[1.35 * inch, 2.65 * inch, 1.15 * inch],
        repeatRows=1,
    )
    prediction_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#287D4A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(prediction_table)
    story.append(Spacer(1, 16))

    if info:
        story.append(Paragraph("Disease Information", styles["Heading1"]))
        story.append(Paragraph(
            str(info.get("common_name", "Information unavailable")),
            styles["Heading2"],
        ))
        story.append(Paragraph(
            str(info.get("overview", "")),
            styles["BodyText"],
        ))
        story.append(Spacer(1, 8))

        for heading, key in [
            ("Symptoms", "symptoms"),
            ("Prevention and Management", "prevention"),
        ]:
            items = info.get(key, [])
            if items:
                story.append(Paragraph(heading, styles["Heading2"]))
                for item in items:
                    story.append(Paragraph(
                        f"&bull; {item}",
                        styles["BodyText"],
                    ))
                    story.append(Spacer(1, 3))
                story.append(Spacer(1, 8))

        story.append(Paragraph("Treatment Guidance", styles["Heading2"]))
        story.append(Paragraph(
            str(info.get("treatment_note", "No treatment guidance available.")),
            styles["BodyText"],
        ))
        story.append(Spacer(1, 10))

        sources = info.get("sources", [])
        if sources:
            story.append(Paragraph("Sources and References", styles["Heading2"]))
            for source in sources:
                title = str(source.get("title", "Reference"))
                url = str(source.get("url", ""))
                safe_url = (
                    url.replace("&", "&amp;")
                    .replace("<", "&lt;")
                    .replace(">", "&gt;")
                )
                story.append(Paragraph(
                    f"{title}: {safe_url}",
                    styles["BodyText"],
                ))
                story.append(Spacer(1, 4))

    story.append(Spacer(1, 14))
    story.append(Paragraph("Important Disclaimer", styles["Heading1"]))
    story.append(Paragraph(
        "This report contains AI-generated estimates for educational "
        "demonstration. Model scores are not confirmed diagnostic "
        "probabilities. Similar symptoms can have different causes. "
        "Consult a qualified plant-disease specialist before taking "
        "treatment action. Do not choose or apply pesticides based only "
        "on this report; follow local regulations and the product label.",
        styles["BodyText"],
    ))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()